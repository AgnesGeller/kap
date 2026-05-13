from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any
import json
import sys

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PROJECT_ROOT / "src" / "data" / "budgetWorkbookData.json"


def clean_value(value: Any) -> Any:
  if value is None:
    return None
  if isinstance(value, (datetime, date, time)):
    return value.isoformat()
  if isinstance(value, Decimal):
    return float(value)
  return value


def normalize_header(value: Any, fallback: str) -> str:
  if value is None or str(value).strip() == "":
    return fallback
  return str(value).strip()


def row_has_value(row: list[Any]) -> bool:
  return any(value is not None and str(value).strip() != "" for value in row)


def extract_table(ws: Any, table: Any) -> dict[str, Any]:
  min_col, min_row, max_col, max_row = range_boundaries(table.ref)
  header_cells = next(
    ws.iter_rows(
      min_row=min_row,
      max_row=min_row,
      min_col=min_col,
      max_col=max_col,
      values_only=True,
    )
  )
  headers = [
    normalize_header(value, f"column_{index + 1}")
    for index, value in enumerate(header_cells)
  ]
  rows: list[dict[str, Any]] = []

  for excel_row_index, row in enumerate(
    ws.iter_rows(
      min_row=min_row + 1,
      max_row=max_row,
      min_col=min_col,
      max_col=max_col,
      values_only=True,
    ),
    start=min_row + 1,
  ):
    values = [clean_value(value) for value in row]
    if not row_has_value(values):
      continue

    rows.append(
      {
        "_excelRow": excel_row_index,
        **{headers[index]: values[index] for index in range(len(headers))},
      }
    )

  return {
    "name": table.name,
    "range": table.ref,
    "headers": headers,
    "rows": rows,
  }


def extract_sheet_rows(ws: Any) -> list[dict[str, Any]]:
  rows: list[dict[str, Any]] = []
  for row in ws.iter_rows():
    values = [clean_value(cell.value) for cell in row]
    if not row_has_value(values):
      continue
    rows.append(
      {
        "row": row[0].row,
        "values": values,
      }
    )
  return rows


def extract_formulas(ws: Any) -> list[dict[str, str]]:
  formulas: list[dict[str, str]] = []
  for row in ws.iter_rows():
    for cell in row:
      if isinstance(cell.value, str) and cell.value.startswith("="):
        formulas.append({"cell": cell.coordinate, "formula": cell.value})
  return formulas


def main() -> None:
  if len(sys.argv) > 1:
    workbook_path = Path(sys.argv[1])
  else:
    workbook_path = Path.home() / "Downloads" / "Költségvetés 2026 (1).xlsx"

  if not workbook_path.exists():
    raise SystemExit(f"Workbook not found: {workbook_path}")

  formula_wb = load_workbook(workbook_path, data_only=False, read_only=False)
  value_wb = load_workbook(workbook_path, data_only=True, read_only=False)

  sheets: dict[str, Any] = {}
  for formula_ws in formula_wb.worksheets:
    value_ws = value_wb[formula_ws.title]
    sheets[formula_ws.title] = {
      "title": formula_ws.title,
      "rowCount": formula_ws.max_row,
      "columnCount": formula_ws.max_column,
      "tables": [
        extract_table(value_ws, table)
        for table in formula_ws.tables.values()
      ],
      "rows": extract_sheet_rows(value_ws),
      "formulas": extract_formulas(formula_ws),
    }

  output = {
    "sourceFile": str(workbook_path),
    "generatedAt": datetime.now().isoformat(timespec="seconds"),
    "sheets": sheets,
  }

  DEFAULT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
  DEFAULT_OUTPUT.write_text(
    json.dumps(output, ensure_ascii=False, indent=2),
    encoding="utf-8",
  )
  print(DEFAULT_OUTPUT)


if __name__ == "__main__":
  main()
